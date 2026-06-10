"""
报告构建器模块 - 项目报告生成

该模块负责构建和输出项目报告，支持 Markdown 和 PPT 两种格式。

核心功能：
- 生成 Markdown 格式报告（包含项目背景、数据集分析、算法原理等章节）
- 生成 PPT 格式报告（支持 python-pptx 库或纯 XML 生成）
- 可自定义报告标题和实验指标

使用示例：
    builder = ReportBuilder()
    
    # 添加指标
    metrics = {"准确率": "95%", "召回率": "92%", "mAP": "0.93"}
    
    # 生成 Markdown 报告
    builder.write_markdown(Path("outputs/report.md"), metrics=metrics)
    
    # 生成 PPT 报告
    builder.write_pptx(Path("outputs/report.pptx"), metrics=metrics)
"""

from dataclasses import dataclass
from pathlib import Path
from typing import Mapping
from xml.sax.saxutils import escape
import zipfile


@dataclass
class ReportBuilder:
    """
    报告构建器，用于生成项目报告
    
    属性：
        default_title: 默认报告标题
    
    方法：
        build_markdown(): 构建 Markdown 报告内容
        write_markdown(): 写入 Markdown 文件
        write_pptx(): 写入 PPT 文件
        _write_minimal_pptx(): 无依赖生成 PPT（纯 XML）
    """

    default_title: str = "3D打印机喷头检测项目报告"  # 默认报告标题

    def build_markdown(self, title: str | None = None, metrics: Mapping[str, object] | None = None) -> str:
        """
        构建 Markdown 格式报告内容
        
        参数：
            title: 报告标题，默认为 default_title
            metrics: 实验指标字典，如 {"准确率": "95%", "召回率": "92%"}
        
        返回值：
            str: Markdown 格式的报告内容
        
        报告结构：
            1. 项目背景：项目目的和应用场景
            2. 数据集分析：数据集来源和预处理流程
            3. 算法原理：YOLOv5 目标检测框架介绍
            4. 模型设计：工厂模式架构说明
            5. 实验过程：实验记录内容说明
            6. 结果分析：实验指标展示
            7. 结论与展望：后续优化方向
        
        示例：
            builder = ReportBuilder()
            metrics = {"准确率": "95%", "mAP": "0.93"}
            md = builder.build_markdown(title="我的报告", metrics=metrics)
        """
        # 使用自定义标题或默认标题
        report_title = title or self.default_title
        
        # 使用提供的指标或空字典
        metrics = metrics or {}
        
        # 将指标转换为 Markdown 列表格式
        metric_lines = "\n".join(f"- `{key}`: {value}" for key, value in metrics.items()) \
            or "- 暂无训练指标，等待实验完成后补充。"
        
        # 返回完整的 Markdown 报告内容
        return f"""# {report_title}

## 项目背景

本项目面向消费级 3D 打印机场景，识别喷头 OK/NG 状态，为打印前报警和停机策略提供算法依据。

## 数据集分析

数据集由外层 `dataset_2` 提供，训练前合并原始训练集和验证集，完成去重、NG/OK 标签归并和分层重划分。

## 算法原理

项目基于 YOLOv5 目标检测框架，通过两类检测头输出 `NG` 和 `OK`。业务上重点关注 NG 类精确率、召回率和低置信度样本。

## 模型设计

项目层使用工厂模式封装数据处理、模型配置、训练、评估和报告生成，YOLOv5 核心源码保持相对稳定。

## 实验过程

实验过程记录数据版本、清洗报告、增强策略、训练命令、权重路径和验证指标。

## 结果分析

{metric_lines}

## 结论与展望

后续应结合误检、漏检和低置信度样本继续修订标注规范、增强困难样本，并评估端侧推理耗时。
"""

    def write_markdown(self, output_path: Path, metrics: Mapping[str, object] | None = None, title: str | None = None) -> Path:
        """
        将 Markdown 报告写入文件
        
        参数：
            output_path: 输出文件路径
            metrics: 实验指标字典
            title: 报告标题
        
        返回值：
            Path: 实际输出的文件路径
        
        执行步骤：
            1. 确保输出目录存在
            2. 构建 Markdown 内容
            3. 写入文件
        
        示例：
            builder = ReportBuilder()
            builder.write_markdown(
                Path("outputs/report.md"),
                metrics={"准确率": "95%"},
                title="喷嘴检测报告"
            )
        """
        # 确保输出目录存在
        output_path.parent.mkdir(parents=True, exist_ok=True)
        
        # 构建并写入 Markdown 内容
        output_path.write_text(self.build_markdown(title=title, metrics=metrics), encoding="utf-8")
        
        return output_path

    def write_pptx(self, output_path: Path, metrics: Mapping[str, object] | None = None, title: str | None = None) -> Path | None:
        """
        将报告写入 PPT 文件
        
        参数：
            output_path: 输出文件路径
            metrics: 实验指标字典
            title: 报告标题
        
        返回值：
            Path: 实际输出的文件路径
        
        实现方式：
            - 如果安装了 python-pptx 库，使用该库生成 PPT
            - 如果未安装，使用纯 XML 方式生成最小化 PPT
        
        PPT 结构：
            1. 封面页：报告标题
            2. 项目目标：介绍项目目标
            3. 实验指标：展示实验结果
            4. 后续优化：列出改进方向
        
        示例：
            builder = ReportBuilder()
            builder.write_pptx(
                Path("outputs/report.pptx"),
                metrics={"准确率": "95%", "mAP": "0.93"}
            )
        """
        try:
            # 尝试导入 python-pptx 库
            from pptx import Presentation
        except ImportError:
            # 如果未安装，使用纯 XML 方式生成
            return self._write_minimal_pptx(output_path, metrics=metrics, title=title)

        # 使用 python-pptx 库生成 PPT
        report_title = title or self.default_title
        prs = Presentation()
        
        # 添加封面页
        self._add_title_slide(prs, report_title)
        
        # 添加内容页
        self._add_bullet_slide(prs, "项目目标", [
            "两类检测：NG / OK", 
            "优先保障高精确率", 
            "保留 YOLOv5 核心能力，新增工厂模式项目层"
        ])
        self._add_bullet_slide(prs, "实验指标", 
            [f"{key}: {value}" for key, value in (metrics or {}).items()] or ["暂无训练指标"]
        )
        self._add_bullet_slide(prs, "后续优化", [
            "复查误检漏检样本", 
            "补充困难样本增强", 
            "验证端侧推理耗时"
        ])
        
        # 确保输出目录存在并保存
        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_path)
        
        return output_path

    def _write_minimal_pptx(self, output_path: Path, metrics: Mapping[str, object] | None = None, title: str | None = None) -> Path:
        """
        无依赖生成最小化 PPT 文件（纯 XML 方式）
        
        当未安装 python-pptx 库时，使用纯 Python 的 zipfile 和 XML 生成 PPT。
        PPTX 本质上是一个 ZIP 文件，包含多个 XML 文件。
        
        参数：
            output_path: 输出文件路径
            metrics: 实验指标字典
            title: 报告标题
        
        返回值：
            Path: 实际输出的文件路径
        
        PPTX 文件结构：
            - [Content_Types].xml: 内容类型定义
            - _rels/.rels: 根关系文件
            - ppt/presentation.xml: 演示文稿主文件
            - ppt/_rels/presentation.xml.rels: 演示文稿关系
            - ppt/slides/slide1.xml, slide2.xml, ...: 幻灯片内容
        
        示例：
            # 当 python-pptx 未安装时自动调用
            builder.write_pptx(Path("outputs/report.pptx"))
        """
        report_title = title or self.default_title
        
        # 定义幻灯片内容
        slides = [
            (report_title, ["YOLOv5 工厂模式项目汇报"]),
            ("项目目标", ["两类检测：NG / OK", "优先保障高精确率", "外层数据不纳入 Git 管理"]),
            ("实验指标", [f"{key}: {value}" for key, value in (metrics or {}).items()] or ["暂无训练指标"]),
            ("后续优化", ["复查误检漏检样本", "补充困难样本增强", "验证端侧推理耗时"]),
        ]
        
        # 确保输出目录存在
        output_path.parent.mkdir(parents=True, exist_ok=True)
        
        # 创建 ZIP 文件（PPTX 格式本质是 ZIP）
        with zipfile.ZipFile(output_path, "w", compression=zipfile.ZIP_DEFLATED) as pptx:
            pptx.writestr("[Content_Types].xml", _content_types(len(slides)))
            pptx.writestr("_rels/.rels", _root_rels())
            pptx.writestr("ppt/presentation.xml", _presentation_xml(len(slides)))
            pptx.writestr("ppt/_rels/presentation.xml.rels", _presentation_rels(len(slides)))
            for index, (slide_title, bullets) in enumerate(slides, start=1):
                pptx.writestr(f"ppt/slides/slide{index}.xml", _slide_xml(slide_title, bullets))
        
        return output_path

    @staticmethod
    def _add_title_slide(prs, title: str) -> None:
        """
        使用 python-pptx 添加封面页幻灯片
        
        参数：
            prs: Presentation 对象
            title: 报告标题
        """
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = title
        slide.placeholders[1].text = "YOLOv5 工厂模式项目汇报"

    @staticmethod
    def _add_bullet_slide(prs, title: str, bullets: list[str]) -> None:
        """
        使用 python-pptx 添加带项目符号的内容页幻灯片
        
        参数：
            prs: Presentation 对象
            title: 幻灯片标题
            bullets: 项目符号列表
        """
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.clear()
        for index, bullet in enumerate(bullets):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0


def _content_types(slide_count: int) -> str:
    """
    生成 [Content_Types].xml 文件内容
    
    PPTX 文件必须包含此文件，定义包中各部分的内容类型。
    
    参数：
        slide_count: 幻灯片数量
    
    返回值：
        str: XML 格式的内容类型定义
    """
    # 为每个幻灯片生成内容类型覆盖
    slide_overrides = "\n".join(
        f'<Override PartName="/ppt/slides/slide{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
        for i in range(1, slide_count + 1)
    )
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>
  {slide_overrides}
</Types>'''


def _root_rels() -> str:
    """
    生成根目录的 _rels/.rels 文件内容
    
    定义包中主要部件之间的关系。
    
    返回值：
        str: XML 格式的关系定义
    """
    return '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>
</Relationships>'''


def _presentation_xml(slide_count: int) -> str:
    """
    生成 ppt/presentation.xml 文件内容
    
    定义演示文稿的主结构，包括幻灯片大小和幻灯片列表。
    
    参数：
        slide_count: 幻灯片数量
    
    返回值：
        str: XML 格式的演示文稿定义
    """
    # 生成幻灯片 ID 列表
    slide_ids = "\n".join(f'<p:sldId id="{255 + i}" r:id="rId{i}"/>' for i in range(1, slide_count + 1))
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:sldSz cx="12192000" cy="6858000" type="wide"/>  <!-- 宽屏格式 -->
  <p:sldIdLst>{slide_ids}</p:sldIdLst>
</p:presentation>'''


def _presentation_rels(slide_count: int) -> str:
    """
    生成 ppt/_rels/presentation.xml.rels 文件内容
    
    定义演示文稿与各幻灯片之间的关系。
    
    参数：
        slide_count: 幻灯片数量
    
    返回值：
        str: XML 格式的关系定义
    """
    rels = "\n".join(
        f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide{i}.xml"/>'
        for i in range(1, slide_count + 1)
    )
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">{rels}</Relationships>'''


def _slide_xml(title: str, bullets: list[str]) -> str:
    """
    生成单个幻灯片的 XML 内容
    
    参数：
        title: 幻灯片标题
        bullets: 项目符号列表
    
    返回值：
        str: XML 格式的幻灯片内容
    
    幻灯片结构：
        - 标题区域：显示标题文字
        - 内容区域：显示项目符号列表
    """
    # 将项目符号转换为 XML 格式
    bullet_text = "".join(
        f'<a:p><a:r><a:rPr lang="zh-CN" sz="2400"/><a:t>{escape(bullet)}</a:t></a:r></a:p>'
        for bullet in bullets
    )
    return f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:cSld>
    <p:spTree>
      <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
      <p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>
      <!-- 标题形状 -->
      <p:sp>
        <p:nvSpPr><p:cNvPr id="2" name="Title"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
        <p:spPr><a:xfrm><a:off x="685800" y="457200"/><a:ext cx="10668000" cy="900000"/></a:xfrm></p:spPr>
        <p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="zh-CN" sz="3600" b="1"/><a:t>{escape(title)}</a:t></a:r></a:p></p:txBody>
      </p:sp>
      <!-- 内容形状（项目符号） -->
      <p:sp>
        <p:nvSpPr><p:cNvPr id="3" name="Content"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
        <p:spPr><a:xfrm><a:off x="914400" y="1600200"/><a:ext cx="10058400" cy="4267200"/></a:xfrm></p:spPr>
        <p:txBody><a:bodyPr/><a:lstStyle/>{bullet_text}</p:txBody>
      </p:sp>
    </p:spTree>
  </p:cSld>
</p:sld>'''
